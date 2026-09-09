# Copyright (c) 2026 Tevinch. SPDX-License-Identifier: MIT
"""Read a column of literal identifiers with an explicit per-record report."""
from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime, time
import io
from numbers import Number
from pathlib import Path
import zipfile

MAX_BYTES = 5_000_000
MAX_ROWS = 10_000  # Data records; the header is separate.
MAX_COLUMNS = 256
MAX_CELL_CHARS = 100_000
MAX_XLSX_EXPANDED_BYTES = 50_000_000


class ImportFailure(ValueError):
    """A fixed error code, optionally with a one-based table record number."""

    def __init__(self, code: str, row: int | None = None):
        self.code, self.row = code, row
        super().__init__(code + (f' at record {row}' if row is not None else ''))


@dataclass(frozen=True)
class Identifier:
    row: int
    value: str


@dataclass(frozen=True)
class RowNotice:
    row: int
    reason: str


@dataclass(frozen=True)
class ImportReport:
    accepted: tuple[Identifier, ...]
    skipped: tuple[RowNotice, ...]
    rejected: tuple[RowNotice, ...]

    @property
    def values(self) -> tuple[str, ...]:
        return tuple(item.value for item in self.accepted)

    @property
    def total_rows(self) -> int:
        return len(self.accepted) + len(self.skipped) + len(self.rejected)


def _validate_column(column):
    if not isinstance(column, str) or not column.strip():
        raise TypeError('Expected a nonblank column name')


def _read_bytes(path):
    try:
        with Path(path).open('rb') as source:
            data = source.read(MAX_BYTES + 1)
    except OSError:
        raise ImportFailure('READ_ERROR') from None
    if len(data) > MAX_BYTES:
        raise ImportFailure('MAX_BYTES')
    return data


def _check_dimensions(values, row):
    if len(values) > MAX_COLUMNS:
        raise ImportFailure('MAX_COLUMNS', row)
    if any(isinstance(value, str) and len(value) > MAX_CELL_CHARS for value in values):
        raise ImportFailure('MAX_CELL_CHARS', row)


def _header_index(header, column):
    if header is None:
        raise ImportFailure('EMPTY_TABLE')
    _check_dimensions(header, 1)
    if not header or any(not isinstance(name, str) or not name.strip() for name in header):
        raise ImportFailure('INVALID_HEADER', 1)
    if len(set(header)) != len(header):
        raise ImportFailure('DUPLICATE_HEADER', 1)
    if column not in header:
        raise ImportFailure('COLUMN_NOT_FOUND', 1)
    return header.index(column)


def _collect(records, index, width):
    accepted, skipped, rejected = [], [], []
    for row, (values, kinds) in enumerate(records, start=2):
        if row > MAX_ROWS + 1:
            raise ImportFailure('MAX_ROWS', row)
        _check_dimensions(values, row)
        if not values:  # An explicit blank CSV record.
            skipped.append(RowNotice(row, 'empty'))
            continue
        if len(values) != width:
            rejected.append(RowNotice(row, 'missing_fields' if len(values) < width else 'extra_fields'))
            continue
        value, kind = values[index], kinds[index]
        if value is None or value == '':
            skipped.append(RowNotice(row, 'empty'))
        elif kind != 'text':
            rejected.append(RowNotice(row, kind))
        elif not value.strip():
            skipped.append(RowNotice(row, 'blank'))
        else:
            accepted.append(Identifier(row, value))
    return ImportReport(tuple(accepted), tuple(skipped), tuple(rejected))


def read_csv_identifiers(path, column: str, *, delimiter: str = ',') -> ImportReport:
    """UTF-8 CSV/TSV. Row numbers count CSV records, including the header."""
    _validate_column(column)
    if delimiter not in (',', ';', '\t'):
        raise TypeError('Choose comma, semicolon or tab')
    try:
        text = _read_bytes(path).decode('utf-8-sig')
    except UnicodeDecodeError:
        raise ImportFailure('INVALID_ENCODING') from None
    reader = csv.reader(io.StringIO(text, newline=''), delimiter=delimiter, strict=True)
    try:
        header = next(reader, None)
        index = _header_index(header, column)
        return _collect(((values, ['text'] * len(values)) for values in reader), index, len(header))
    except csv.Error:
        raise ImportFailure('INVALID_CSV') from None


def _cell_kind(cell):
    if cell.data_type == 'f':
        return 'formula'
    if cell.data_type == 'e':
        return 'error'
    value = cell.value
    if value is None:
        return 'empty'
    if isinstance(value, bool):
        return 'boolean'
    if isinstance(value, (datetime, date, time)):
        return 'date'
    if isinstance(value, Number):
        return 'numeric'
    if isinstance(value, str):
        return 'text'
    return 'nontext'


def read_xlsx_identifiers(path, column: str, *, sheet_name: str) -> ImportReport:
    """Read stored text cells; never infer identifiers from numbers or formats."""
    _validate_column(column)
    if not isinstance(sheet_name, str) or not sheet_name:
        raise TypeError('Expected an explicit sheet name')
    data = _read_bytes(path)
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(part.file_size for part in archive.infolist()) > MAX_XLSX_EXPANDED_BYTES:
                raise ImportFailure('MAX_XLSX_EXPANDED_BYTES')
    except zipfile.BadZipFile:
        raise ImportFailure('INVALID_XLSX') from None
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportFailure('XLSX_DEPENDENCY') from None
    workbook = None
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        if sheet_name not in workbook.sheetnames:
            raise ImportFailure('SHEET_NOT_FOUND')
        sheet = workbook[sheet_name]
        if (sheet.max_row or 0) > MAX_ROWS + 1:
            raise ImportFailure('MAX_ROWS')
        declared_columns = sheet.max_column or 0
        if declared_columns > MAX_COLUMNS:
            raise ImportFailure('MAX_COLUMNS')
        # An understated dimension must not truncate stored rows or columns.
        sheet.reset_dimensions()
        rows = sheet.iter_rows()
        first = next(rows, None)
        header = None if first is None else [cell.value for cell in first]
        if header is not None:
            header += [None] * max(0, declared_columns - len(header))
        if first is not None and any(_cell_kind(cell) != 'text' for cell in first):
            raise ImportFailure('INVALID_HEADER', 1)
        index = _header_index(header, column)
        def records():
            for row in rows:
                values = [cell.value for cell in row]
                kinds = [_cell_kind(cell) for cell in row]
                # XLSX omits empty cells; missing trailing cells are still empty.
                padding = max(0, len(header) - len(values))
                yield values + [None] * padding, kinds + ['empty'] * padding

        return _collect(records(), index, len(header))
    except ImportFailure:
        raise
    except Exception:
        # Library errors may contain workbook values; expose only a fixed code.
        raise ImportFailure('INVALID_XLSX') from None
    finally:
        if workbook is not None:
            workbook.close()
