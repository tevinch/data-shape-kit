# Copyright (c) 2026 Tevinch. SPDX-License-Identifier: MIT
import csv
from datetime import datetime
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from identifier_import import ImportFailure, read_csv_identifiers, read_xlsx_identifiers

try:
    import openpyxl
except ImportError:
    openpyxl = None


class CsvTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'sample.csv'

    def load(self, text, column='ID', **kwargs):
        self.path.write_bytes(text.encode('utf-8'))
        return read_csv_identifiers(self.path, column, **kwargs)

    def failure(self, text, code, **kwargs):
        with self.assertRaises(ImportFailure) as error:
            self.load(text, **kwargs)
        self.assertEqual(error.exception.code, code)
        return error.exception

    def test_original_three_identifiers_survive(self):
        report = self.load('ID\n00123\n00007\nNA\n')
        self.assertEqual(report.values, ('00123', '00007', 'NA'))
        self.assertEqual(report.total_rows, 3)
        self.assertEqual(report.skipped, ())
        self.assertEqual(report.rejected, ())

    def test_literals_long_identifiers_unicode_and_spaces_remain_exact(self):
        values = ['NULL', 'N/A', 'None', 'nan', '123456789012345678901234567890', '零件-００１', '  001  ', '=A1']
        text = io.StringIO(newline='')
        writer = csv.writer(text); writer.writerow(['ID']); writer.writerows([[x] for x in values])
        self.assertEqual(self.load(text.getvalue()).values, tuple(values))

    def test_blank_empty_missing_and_extra_records_are_accounted_for(self):
        report = self.load('Other,ID\na,001\nb,\nc,   \nd\ne,002,extra\n\nf,003\n')
        self.assertEqual(report.values, ('001', '003'))
        self.assertEqual([(x.row, x.reason) for x in report.skipped], [(3, 'empty'), (4, 'blank'), (7, 'empty')])
        self.assertEqual([(x.row, x.reason) for x in report.rejected], [(5, 'missing_fields'), (6, 'extra_fields')])
        self.assertEqual(report.total_rows, 7)

    def test_embedded_newlines_use_record_numbers_not_physical_lines(self):
        report = self.load('ID,Note\r\n"00\r\n1",ok\r\n002,\r\n')
        self.assertEqual(report.values, ('00\r\n1', '002'))
        self.assertEqual([x.row for x in report.accepted], [2, 3])

    def test_bom_and_explicit_delimiters(self):
        for delimiter in [',', ';', '\t']:
            with self.subTest(delimiter=delimiter):
                report = self.load(f'\ufeffName{delimiter}ID\npart{delimiter}001\n', delimiter=delimiter)
                self.assertEqual(report.values, ('001',))

    def test_duplicate_identifiers_are_not_deduplicated(self):
        self.assertEqual(self.load('ID\n001\n001\n').values, ('001', '001'))

    def test_header_match_is_exact_and_case_sensitive(self):
        self.assertEqual(self.load(' ID ,id\n001,002\n', column=' ID ').values, ('001',))
        self.failure(' ID ,id\n001,002\n', 'COLUMN_NOT_FOUND')

    def test_duplicate_or_blank_headers_fail(self):
        self.failure('ID,ID\n1,2\n', 'DUPLICATE_HEADER')
        self.failure('ID, \n1,2\n', 'INVALID_HEADER')

    def test_no_header_and_header_only(self):
        self.failure('', 'EMPTY_TABLE')
        self.failure('\ufeff', 'EMPTY_TABLE')
        self.assertEqual(self.load('ID\n').total_rows, 0)

    def test_invalid_encoding_and_csv_fail_without_input_values(self):
        self.path.write_bytes(b'ID\n\xff\n')
        with self.assertRaises(ImportFailure) as caught:
            read_csv_identifiers(self.path, 'ID')
        self.assertEqual(caught.exception.code, 'INVALID_ENCODING')
        error = self.failure('ID\n"PRIVATE-INCOMPLETE', 'INVALID_CSV')
        self.assertNotIn('PRIVATE', str(error))

    def test_bad_options_fail(self):
        for column, delimiter in [('', ','), ('  ', ','), ('ID', '|'), (1, ',')]:
            with self.subTest(column=column, delimiter=delimiter):
                with self.assertRaises(TypeError):
                    self.load('ID\n001\n', column=column, delimiter=delimiter)

    def test_size_row_column_and_cell_limits_abort_instead_of_truncating(self):
        self.failure('ID\n' + 'a' * 5_000_000, 'MAX_BYTES')
        self.failure('ID\n' + 'x\n' * 10_001, 'MAX_ROWS')
        self.failure(','.join(f'H{i}' for i in range(257)), 'MAX_COLUMNS', column='H0')
        self.failure('ID\n' + 'a' * 100_001, 'MAX_CELL_CHARS')

    def test_exact_row_limit_and_cell_limit_are_accepted(self):
        self.assertEqual(self.load('ID\n' + 'x\n' * 10_000).total_rows, 10_000)
        self.assertEqual(len(self.load('ID\n' + 'a' * 100_000).values[0]), 100_000)

    def test_csv_path_does_not_import_pandas_or_openpyxl(self):
        with patch.dict('sys.modules', {'pandas': None, 'openpyxl': None}):
            self.assertEqual(self.load('ID\n001\nNA\n').values, ('001', 'NA'))

    def test_source_file_is_not_modified(self):
        content = b'ID\r\n001\r\nNA\r\n'
        self.path.write_bytes(content)
        read_csv_identifiers(self.path, 'ID')
        self.assertEqual(self.path.read_bytes(), content)

    def test_read_error_does_not_echo_the_path(self):
        with self.assertRaises(ImportFailure) as caught:
            read_csv_identifiers(Path(self.temp.name) / 'private-name.csv', 'ID')
        self.assertEqual(caught.exception.code, 'READ_ERROR')
        self.assertNotIn('private-name', str(caught.exception))


@unittest.skipUnless(openpyxl, 'Optional XLSX checks need openpyxl')
class XlsxTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'sample.xlsx'

    def workbook(self, rows):
        workbook = openpyxl.Workbook()
        sheet = workbook.active; sheet.title = 'Items'
        for row in rows: sheet.append(row)
        return workbook, sheet

    def save_load(self, workbook, **kwargs):
        workbook.save(self.path); workbook.close()
        return read_xlsx_identifiers(self.path, 'ID', sheet_name='Items', **kwargs)

    def test_same_text_contract_as_csv(self):
        values = ['00123', '00007', 'NA', 'NULL', '123456789012345678901234', '零件-００１', ' 001 ', 'a\nb']
        workbook, _ = self.workbook([['ID'], *[[x] for x in values]])
        report = self.save_load(workbook)
        self.assertEqual(report.values, tuple(values))
        self.assertEqual(report.total_rows, len(values))

    def test_numeric_formats_dates_booleans_formulas_and_errors_are_rejected(self):
        workbook, sheet = self.workbook([['ID'], ['00123'], [123], [True], [datetime(2026, 9, 9)], ['=1+1'], ['#N/A'], ['NULL']])
        sheet['A3'].number_format = '00000'
        report = self.save_load(workbook)
        self.assertEqual(report.values, ('00123', 'NULL'))
        self.assertEqual([(x.row, x.reason) for x in report.rejected], [(3, 'numeric'), (4, 'boolean'), (5, 'date'), (6, 'formula'), (7, 'error')])
        self.assertEqual(report.total_rows, 7)

    def test_formula_looking_text_is_distinct_from_a_formula_cell(self):
        workbook, sheet = self.workbook([['ID'], ['=A1'], ['=A1']])
        sheet['A2'].data_type = 's'
        report = self.save_load(workbook)
        self.assertEqual(report.values, ('=A1',))
        self.assertEqual([(x.row, x.reason) for x in report.rejected], [(3, 'formula')])

    def test_empty_and_whitespace_rows_are_counted(self):
        workbook, _ = self.workbook([['ID', 'Note'], ['001', 'a'], [None, 'b'], ['  ', 'c'], [None, None], ['002', 'd']])
        report = self.save_load(workbook)
        self.assertEqual(report.values, ('001', '002'))
        self.assertEqual([(x.row, x.reason) for x in report.skipped], [(3, 'empty'), (4, 'blank'), (5, 'empty')])

    def test_selected_sheet_is_explicit(self):
        workbook, _ = self.workbook([['ID'], ['001']]); workbook.save(self.path); workbook.close()
        with self.assertRaises(ImportFailure) as caught:
            read_xlsx_identifiers(self.path, 'ID', sheet_name='Missing')
        self.assertEqual(caught.exception.code, 'SHEET_NOT_FOUND')
        with self.assertRaises(TypeError):
            read_xlsx_identifiers(self.path, 'ID', sheet_name='')

    def test_invalid_header_and_merged_range_are_not_guessed(self):
        workbook, sheet = self.workbook([['ID', 'Note'], ['001', 'x']])
        sheet.merge_cells('A1:B1')
        with self.assertRaises(ImportFailure) as caught: self.save_load(workbook)
        self.assertEqual(caught.exception.code, 'INVALID_HEADER')

    def test_sheet_dimensions_are_bounded(self):
        workbook, sheet = self.workbook([['ID'], ['001']]); sheet.cell(10002, 1, 'x')
        with self.assertRaises(ImportFailure) as caught: self.save_load(workbook)
        self.assertEqual(caught.exception.code, 'MAX_ROWS')

    def test_underreported_sheet_dimensions_do_not_hide_rows_or_columns(self):
        workbook, _ = self.workbook([['Note', 'ID'], ['first', '001'], ['second', 'NA'], ['third', None]])
        workbook.save(self.path); workbook.close()
        buffer = io.BytesIO()
        with zipfile.ZipFile(self.path) as source, zipfile.ZipFile(buffer, 'w') as target:
            for name in source.namelist():
                data = source.read(name)
                if name == 'xl/worksheets/sheet1.xml':
                    data = data.replace(b'<dimension ref="A1:B4"/>', b'<dimension ref="A1:A1"/>')
                target.writestr(name, data)
        self.path.write_bytes(buffer.getvalue())
        report = read_xlsx_identifiers(self.path, 'ID', sheet_name='Items')
        self.assertEqual(report.values, ('001', 'NA'))
        self.assertEqual([(x.row, x.reason) for x in report.skipped], [(4, 'empty')])
        self.assertEqual(report.total_rows, 3)

    def test_invalid_archive_and_missing_dependency_have_fixed_errors(self):
        self.path.write_bytes(b'not an xlsx')
        with self.assertRaises(ImportFailure) as caught:
            read_xlsx_identifiers(self.path, 'ID', sheet_name='Items')
        self.assertEqual(caught.exception.code, 'INVALID_XLSX')
        workbook, _ = self.workbook([['ID'], ['001']]); workbook.save(self.path); workbook.close()
        with patch.dict('sys.modules', {'openpyxl': None}):
            with self.assertRaises(ImportFailure) as caught:
                read_xlsx_identifiers(self.path, 'ID', sheet_name='Items')
        self.assertEqual(caught.exception.code, 'XLSX_DEPENDENCY')

    def test_xlsx_source_is_not_modified(self):
        workbook, _ = self.workbook([['ID'], ['001']]); workbook.save(self.path); workbook.close()
        before = self.path.read_bytes()
        read_xlsx_identifiers(self.path, 'ID', sheet_name='Items')
        self.assertEqual(self.path.read_bytes(), before)


if __name__ == '__main__':
    unittest.main()
